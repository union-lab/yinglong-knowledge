#!/usr/bin/env python3
"""Build the DouDian merchant self-developed backend application document."""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


TEMPLATE = Path("/Users/youlian007/Desktop/【模板】商家自研后台系统功能说明书.docx")
SOURCE_DOC = Path("/Users/youlian007/Documents/Codex/2026-07-14/shi/outputs/（修订1.0）友联巡店系统.docx")
SOURCE_IMAGES = Path("/private/tmp/doudian_application_doc/source_images")
CERT = Path("/Users/youlian007/Desktop/API申请文档集/公司证照图片/友联巡店系统软件著作权.jpg")
ASSET_DIR = Path("/private/tmp/doudian_application_doc/assets")
OUTPUT = Path("/Users/youlian007/Desktop/union-knowledgebase/ylLib/友联外部接口申请/抖店API申请/友联巡店系统_抖店商家自研后台系统功能说明书.docx")

FONT_CN = "等线"
FONT_BODY = "等线"
FONT_CODE = "Menlo"
INK = "222222"
TEAL = "087F74"
TEAL_LIGHT = "E7F4F2"
GRAY = "F2F4F5"
MID_GRAY = "D9DEE2"
WHITE = "FFFFFF"


API_ROWS = [
    ("店铺 API", "/shop/getShopCategory；/address/list；/shop/getExperienceScore", "同步授权店铺类目、地址与经营体验信息，用于店铺主数据和运营看板。"),
    ("素材中心 API", "/material/uploadImageSync；/material/searchMaterial；/material/uploadVideoAsync", "管理商品和内容运营所需图片、视频素材，形成可追溯素材档案。"),
    ("商品 API", "/product/listV2；/product/detail；/sku/detail；/spu/getSpu", "在金蝶商品档案基础上补齐抖店 SPU、SKU、规格、状态和平台标识。"),
    ("订单 API", "/order/searchList；/order/orderDetail", "接入抖店订单并转换为系统平台订单（普通订单），进入库存校验和履约流程。"),
    ("库存 API", "/sku/stockNum；/sku/syncStock；/sku/syncStockBatch", "查询与同步平台 SKU 库存；按平台销售单位与金蝶库存单位映射比例换算。"),
    ("售后退款 API", "/afterSale/List；/afterSale/Detail；/afterSale/operate；/afterSale/fillLogistics；/afterSale/returnGoodsToWareHouseSuccess", "同步退款、退货退款等售后单，完成审核、退货物流、入仓和状态回写。"),
    ("物流发货 API", "/order/logisticsAdd；/order/logisticsAddMultiPack；/order/logisticsEdit；/order/queryOrderLogistics", "金蝶销售出库取得物流单号后回传抖店，并支持多包裹和物流纠错。"),
    ("电子面单 API", "/logistics/newCreateOrder；/logistics/waybillApply；/logistics/updateOrder；/logistics/cancelOrder", "按仓库发货需要申请、更新或取消电子面单，保留面单和包裹对应关系。"),
    ("账单 API", "/order/getShopAccountItem；/order/getSettleBillDetailV3；/order/downloadShopAccountItemFile", "同步店铺动账和结算明细，用于订单、退款、佣金和平台费用核对。"),
    ("精选联盟 API", "/buyin/queryShopAllianceProducts；/buyin/queryShopAllianceOrder；/alliance/getOrderList", "同步联盟推广商品、达人/联盟订单及佣金数据，支持后续内容与分销运营。"),
    ("保险 API", "/order/policy；/order/insurance", "查询订单保单及保险服务信息，关联订单与售后责任核验。"),
    ("商户登录风险 API（风控安全 API）", "/antispam/userLogin", "用户登录时提交必要风险要素，识别异常登录并记录风控结果。"),
    ("飞鸽 API", "/pigeon/workOrderNotify；/pigeon/getStaffOnlineRecordList；/pigeon/updatePushUser", "接收/更新客服工单与客服在线记录，关联订单、售后和服务处理。"),
    ("短视频 API", "/shopVideo/listVideo；/shopVideo/uploadVideo；/shopVideo/publishVideo；/shopVideo/listRelatableProduct", "管理店铺短视频、挂载商品和发布状态，服务内容运营。"),
]

FUNCTION_ROWS = [
    ("会员管理", "包含：用户注册申请、付款提报列表、微信用户列表、注册用户管理、用户等级。会员数据按授权店铺隔离。", "店铺/订单授权数据、精选联盟 API"),
    ("订单管理", "包含：普通订单列表、待外采单列表、服务订单列表、发货单列表、售后单列表、退货单列表。", "订单 API、库存 API、售后退款 API、物流发货 API"),
    ("商品管理", "包含：普通商品列表、金蝶商品分类、商品剂型、店内码映射管理、商品分类、品牌列表、商品评价。", "商品 API、素材中心 API、库存 API"),
    ("财务管理", "负责店铺动账、结算、退款及服务订单支付核对；线下支付仅用于运营支付服务订单费用。", "账单 API、售后退款 API"),
    ("促销中心", "一级功能保留，当前尚未启用，后续按业务计划启用。", "精选联盟 API"),
    ("分销设置", "一级功能保留，当前尚未启用，后续按业务计划启用。", "精选联盟 API"),
    ("运营管理", "包含：公告列表、自定义表单、广告管理、文章管理、库存管理、报表统计。", "素材中心 API、库存 API、精选联盟 API、短视频 API"),
    ("微信公众号", "供店铺运营查阅订单数据、处理待办并支付本人承担的服务订单费用。", "内部聚合数据、微信支付"),
    ("商城设置", "包含：平台设置、页面设计、签到设置、支付设置、商城服务说明、配送方式、门店列表、物流公司、区域设置、打印机设置、消息配置。", "店铺 API、物流发货 API"),
    ("后台管理", "负责账号权限、接口任务、参数配置、操作日志和异常处理记录。", "商户登录风险 API、飞鸽 API"),
]

SCREENSHOTS = [
    ("01_p025.png", "系统登录界面", "账号密码登录后执行商户登录风险校验，异常结果进入审计日志。"),
    ("02_p031.png", "系统工作台", "展示订单、售后、库存和履约待办，数据按授权店铺隔离。"),
    ("08_p059.png", "普通订单列表", "抖店订单接入后转换为平台订单，并保留平台订单号、子订单和店铺标识。"),
    ("09_p060.png", "待外采单列表", "库存不足子订单由人工按需生成待外采单，采购补足后重新校验。"),
    ("10_p061.png", "服务订单列表", "运营责任售后正常退货入仓后，以运营人员为客户生成服务订单。"),
    ("11_p062.png", "发货单列表", "金蝶订单创建成功后逐单生成；每 6 小时获取销售出库物流并回传。"),
    ("12_p063.png", "售后单列表", "同步抖店退款/退货退款申请、审核节点及处理状态。"),
    ("13_p064.png", "退货单列表", "记录退货物流、质检、入仓结果，并与售后单、原订单关联。"),
    ("14_p077.png", "普通商品列表", "金蝶商品为基础，合并抖店 SPU/SKU、规格、平台状态和素材信息。"),
    ("17_p080.png", "店内码映射管理", "维护平台 SKU 销售单位与金蝶库存单位之间的编码映射和换算比例。"),
    ("22_p091.png", "账单及退款核对", "按订单、售后和平台动账核对；服务订单支付与平台退款分开管理。"),
    ("26_p102.png", "促销与分销配置", "当前未启用；精选联盟权限获批后按业务计划启用。"),
    ("28_p111.png", "报表统计", "聚合店铺订单、库存、售后、联盟和账单指标，支持运营查阅。"),
    ("29_p116.png", "微信公众号入口", "店铺运营可查阅订单待办并支付本人承担的服务订单费用。"),
    ("30_p125.png", "商城设置", "保留必要的平台、支付、配送、物流、区域和消息配置。"),
    ("39_p142.png", "后台管理", "记录接口调用、任务执行、关键状态变化和人工处理痕迹。"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name=FONT_BODY, size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    existing = {style.name for style in styles}
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet"):
        if name not in existing:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    normal = styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(5)

    for name, size, color in (("Title", 24, INK), ("Heading 1", 18, TEAL), ("Heading 2", 16, INK), ("Heading 3", 14, TEAL)):
        style = styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    for name in ("Caption",):
        style = styles[name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(14)
        style.font.color.rgb = RGBColor.from_string("555555")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color="666666")


def configure_section(section, cover=False) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = cover


def add_header_footer(section) -> None:
    header_p = section.header.paragraphs[0]
    header_p.text = "友联巡店系统 V1.0  |  抖店开放平台 API 申请材料"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_p.runs:
        set_run_font(run, size=8, color="777777")
    p = section.footer.paragraphs[0]
    add_page_number(p)


def add_para(doc, text="", *, bold_prefix=None, align=None, indent=True, size=None, color=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent and style is None:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, size=size, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=size, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.7 + 0.5 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    r = p.add_run(text)
    set_run_font(r, size=14)
    return p


def add_note(doc, title, text, fill=TEAL_LIGHT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(17.2)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{title}  ")
    set_run_font(r, bold=True, size=14, color=TEAL)
    r = p.add_run(text)
    set_run_font(r, size=14, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, font_size=8.5, header_fill=TEAL):
    font_size = max(font_size, 14)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), MID_GRAY)
        borders.append(tag)
    tbl_pr.append(borders)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, bold=True, size=font_size, color=WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for row_data in rows:
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, bold=(i == 0))
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level):
    return doc.add_heading(text, level=level)


def add_cover(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("抖店开放平台")
    set_run_font(r, FONT_CN, 18, True, TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("商家自研后台系统功能说明书")
    set_run_font(r, FONT_CN, 25, True, INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("友联巡店系统 V1.0")
    set_run_font(r, FONT_CN, 17, True, TEAL)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(r, FONT_BODY, 14, False, TEAL)

    cover_rows = [
        ("申请主体", "泉州市友联医疗器械有限公司"),
        ("系统性质", "商家自研后台系统"),
        ("软件版本", "V1.0"),
        ("软著登记号", "2024SR0357042"),
        ("文档版本", "申请版 1.0"),
        ("编制日期", "2026 年 7 月 14 日"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for k, v in cover_rows:
        row = table.add_row()
        row.cells[0].width = Cm(4.2)
        row.cells[1].width = Cm(9.0)
        set_cell_margins(row.cells[0], top=130, bottom=130)
        set_cell_margins(row.cells[1], top=130, bottom=130)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p0.add_run(k)
        set_run_font(r, bold=True, size=14, color="555555")
        p1 = row.cells[1].paragraphs[0]
        r = p1.add_run(v)
        set_run_font(r, size=14, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("用于抖店开放平台 API 权限申请与商家自研系统能力审核")
    set_run_font(r, size=14, color="666666")
    doc.add_page_break()


def add_contents(doc) -> None:
    add_heading(doc, "目录", 1)
    items = [
        "1 概述",
        "  1.1 名词解释",
        "  1.2 产品概述及目标",
        "  1.3 研发信息及说明",
        "2 调用接口",
        "  2.1 调用接口（API）一览表",
        "  2.2 接口调用与数据安全说明",
        "3 功能详情",
        "  3.1 功能一览表",
        "  3.2 会员管理",
        "  3.3 订单管理",
        "  3.4 商品管理",
        "  3.5 财务管理",
        "  3.6 促销中心",
        "  3.7 分销设置",
        "  3.8 运营管理",
        "  3.9 微信公众号",
        "  3.10 商城设置",
        "  3.11 后台管理",
        "4 系统架构、功能截图",
        "  4.1 系统架构、流程图",
        "  4.2 系统登录主界面截图",
        "  4.3 系统子功能界面截图",
        "  4.4 自研发软件著作权",
        "  4.5 自研发软件源代码信息",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4 if not item.startswith("  ") else 1.1)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item.strip())
        set_run_font(r, size=14, bold=not item.startswith("  "), color=TEAL if not item.startswith("  ") else INK)
    doc.add_page_break()


def add_overview(doc) -> None:
    add_heading(doc, "1 概述", 1)
    add_heading(doc, "1.1 名词解释", 2)
    terms = [
        ("抖店", "抖音电商商家经营平台及其开放平台能力。"),
        ("SPU", "标准产品单元，用于描述同一商品的共性信息。"),
        ("SKU", "库存量单位，用于描述商品规格、销售单位和平台库存。"),
        ("平台订单", "本系统接入抖店订单后形成的内部普通订单，是库存校验、金蝶下单和履约的主单据。"),
        ("子订单", "平台订单中的商品明细行；库存按子订单逐行校验。"),
        ("可发库存", "在金蝶库存基础上扣除占用、冻结等不可发数量后，可用于当前订单履约的库存。"),
        ("待外采单", "订单存在可发库存不足时，由人工按需生成并交采购部门补货的内部单据。"),
        ("服务订单", "运营责任售后完成退货入仓后，以责任运营人员为客户生成的费用结算订单。"),
        ("平台 SKU 销售单位映射", "维护平台 SKU 销售单位与金蝶商品库存单位之间的编码关系和换算比例。"),
    ]
    add_table(doc, ["名词", "说明"], terms, widths=[4.0, 13.0], font_size=9)

    add_heading(doc, "1.2 产品概述及目标", 2)
    add_para(doc, "友联巡店系统 V1.0 是泉州市友联医疗器械有限公司自研的商家后台系统，用于统一管理授权抖店店铺的商品、订单、库存、售后、物流、账单、联盟推广、内容和客服业务，并与公司金蝶系统完成商品、库存、订单和出库物流协同。")
    add_para(doc, "本次申请的核心目标是：通过抖店开放平台 API 获得稳定、合规、可审计的数据通道，减少人工导入和重复操作，保障订单及时履约、库存口径一致、售后过程可追踪，并为店铺运营提供统一工作台。")
    add_bullet(doc, "商品一体化：先以金蝶商品为基础档案，再补齐抖店 SPU/SKU、规格、状态、图片和视频素材。")
    add_bullet(doc, "订单自动履约：订单接入、单位换算、可发库存校验、自动审核、金蝶下单及物流回传形成闭环。")
    add_bullet(doc, "异常可干预：库存不足时暂停下推，人工决定是否生成待外采单，补货后重新校验。")
    add_bullet(doc, "售后责任可核算：退货入仓与退款状态同步，运营责任场景自动生成服务订单。")
    add_bullet(doc, "权限最小化：只申请实际业务所需的 14 类 API；短信 API 不申请，工具 API 原则上不申请。")

    add_heading(doc, "1.3 研发信息及说明", 2)
    rows = [
        ("软件名称", "友联巡店系统 V1.0"),
        ("著作权人/申请主体", "泉州市友联医疗器械有限公司"),
        ("软件著作权", "登记号 2024SR0357042；证书号 软著登字第12760915号"),
        ("研发完成日期", "2022 年 7 月 8 日"),
        ("主要技术", ".NET 9 / ASP.NET 9、MySQL、Windows Server / IIS、Edge / Chrome"),
        ("研发方式", "商家自研；产品、后端、前端、测试和运维岗位协作"),
        ("接口负责人", "由申请主体指定，提交前补充姓名与手机号码"),
        ("生产部署", "应用、数据库分离部署，启用访问控制、日志审计和备份；实际云厂商及规格以提交时生产环境清单为准"),
        ("数据使用边界", "仅处理本公司授权店铺数据，不向未授权第三方披露；密钥、令牌和敏感字段加密保存"),
    ]
    add_table(doc, ["项目", "说明"], rows, widths=[4.2, 12.8], font_size=9)


def add_api_chapter(doc) -> None:
    doc.add_page_break()
    add_heading(doc, "2 调用接口", 1)
    add_heading(doc, "2.1 调用接口（API）一览表", 2)
    add_para(doc, "以下权限均直接对应系统实际功能。基础类 API 用于日常店铺经营闭环；精选联盟、保险、商户登录风险、飞鸽和短视频 API 用于推广、保障、风控、客服和内容运营。")
    rows = [(i + 1, *row) for i, row in enumerate(API_ROWS)]
    add_table(doc, ["序号", "API 类别", "主要接口", "业务用途"], rows, widths=[1.0, 3.3, 5.6, 7.1], font_size=7.7)
    add_note(doc, "不申请范围", "短信 API 不需要、不申请；工具 API 原则上不申请。应用授权和令牌刷新仅作为开放平台调用的必要技术前置，不作为业务能力扩张。")

    add_heading(doc, "2.2 接口调用与数据安全说明", 2)
    safeguards = [
        "认证与签名：调用 openapi-fxg.jinritemai.com，使用平台分配的 app_key、access_token 和 HMAC-SHA256 签名；签名实现采用官方 SDK/调用指南并通过官方调试工具校验。",
        "令牌保护：access_token、refresh_token 和 app_secret 仅在服务端加密保存，不进入前端页面、日志正文或本申请材料。",
        "店铺隔离：接口请求、落库、任务和审计日志均携带授权店铺标识，禁止跨店铺读取或写入。",
        "幂等与重试：以平台订单号、售后编号、物流单号和请求幂等键防止重复建单；限流或瞬时失败采用退避重试。",
        "最小数据：只获取业务闭环所需字段；敏感字段按平台规则脱敏展示，并设置访问角色和留痕。",
        "审计与告警：记录接口类别、方法、店铺、请求时间、返回码、耗时和业务主键，不记录密钥及完整敏感报文。",
    ]
    for item in safeguards:
        add_bullet(doc, item)


def draw_flow(filename, title, lanes, note=None):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 980
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
    title_font = ImageFont.truetype(font_path, 52)
    box_font = ImageFont.truetype(font_path, 31)
    small_font = ImageFont.truetype(font_path, 24)
    draw.text((80, 55), title, font=title_font, fill="#222222")
    draw.line((80, 130, width - 80, 130), fill="#087F74", width=7)
    lane_h = 180
    start_y = 180
    for lane_index, (lane_name, steps) in enumerate(lanes):
        y = start_y + lane_index * lane_h
        draw.rounded_rectangle((70, y, 270, y + 120), radius=12, fill="#087F74")
        bbox = draw.textbbox((0, 0), lane_name, font=box_font)
        draw.text((170 - (bbox[2] - bbox[0]) / 2, y + 60 - (bbox[3] - bbox[1]) / 2), lane_name, font=box_font, fill="white")
        available = width - 370
        gap = 28
        box_w = int((available - gap * (len(steps) - 1)) / len(steps))
        x = 330
        for i, step in enumerate(steps):
            fill = "#E7F4F2" if i % 2 == 0 else "#F2F4F5"
            draw.rounded_rectangle((x, y, x + box_w, y + 120), radius=12, fill=fill, outline="#AEB8BC", width=2)
            lines = []
            current = ""
            for ch in step:
                test = current + ch
                if draw.textlength(test, font=small_font) > box_w - 30 and current:
                    lines.append(current)
                    current = ch
                else:
                    current = test
            if current:
                lines.append(current)
            total_h = len(lines) * 32
            ty = y + (120 - total_h) / 2
            for line in lines:
                tw = draw.textlength(line, font=small_font)
                draw.text((x + (box_w - tw) / 2, ty), line, font=small_font, fill="#222222")
                ty += 32
            if i < len(steps) - 1:
                ax = x + box_w + 4
                ay = y + 60
                draw.line((ax, ay, ax + gap - 10, ay), fill="#087F74", width=5)
                draw.polygon([(ax + gap - 10, ay - 9), (ax + gap, ay), (ax + gap - 10, ay + 9)], fill="#087F74")
            x += box_w + gap
    if note:
        draw.rounded_rectangle((80, height - 125, width - 80, height - 55), radius=12, fill="#FFF5D8")
        draw.text((110, height - 110), note, font=small_font, fill="#5E4A00")
    path = ASSET_DIR / filename
    img.save(path, quality=95)
    return path


def draw_system_architecture(filename):
    """Draw the application architecture using the confirmed level-one menu."""
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 1120
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
    title_font = ImageFont.truetype(font_path, 52)
    lane_font = ImageFont.truetype(font_path, 30)
    box_font = ImageFont.truetype(font_path, 26)
    small_font = ImageFont.truetype(font_path, 23)

    draw.text((80, 48), "系统对接架构", font=title_font, fill="#222222")
    draw.line((80, 125, width - 80, 125), fill="#087F74", width=7)

    def lane_label(text, y, h):
        draw.rounded_rectangle((65, y, 265, y + h), radius=12, fill="#087F74")
        bbox = draw.textbbox((0, 0), text, font=lane_font)
        draw.text((165 - (bbox[2] - bbox[0]) / 2, y + h / 2 - (bbox[3] - bbox[1]) / 2),
                  text, font=lane_font, fill="white")

    def centered_box(x, y, w, h, text, fill="#E7F4F2", outline="#AEB8BC", font=box_font):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=10, fill=fill, outline=outline, width=2)
        bbox = draw.textbbox((0, 0), text, font=font)
        draw.text((x + (w - (bbox[2] - bbox[0])) / 2,
                   y + h / 2 - (bbox[3] - bbox[1]) / 2), text, font=font, fill="#222222")

    # Platform capabilities required by this application.
    lane_label("抖店开放平台", 175, 150)
    platform_boxes = ["店铺/素材/商品", "订单/库存", "售后/物流/面单", "账单/保险", "联盟/风控/飞鸽/视频"]
    x0, gap, box_w = 315, 22, 276
    for i, text in enumerate(platform_boxes):
        centered_box(x0 + i * (box_w + gap), 200, box_w, 100, text,
                     fill="#F2F4F5", font=small_font)

    # Shared API gateway and security boundary.
    draw.line((900, 325, 900, 365), fill="#087F74", width=5)
    draw.polygon([(890, 355), (900, 370), (910, 355)], fill="#087F74")
    centered_box(510, 365, 780, 62, "API 网关 · 授权店铺隔离 · 幂等重试 · 脱敏审计",
                 fill="#FFF5D8", outline="#D8C071", font=small_font)

    # Confirmed level-one functions from the system console.
    lane_label("友联巡店系统", 455, 310)
    modules = [
        "会员管理", "订单管理", "商品管理", "财务管理", "促销中心",
        "分销设置", "运营管理", "微信公众号", "商城设置", "后台管理",
    ]
    module_w, module_h = 270, 92
    for i, text in enumerate(modules):
        row, col = divmod(i, 5)
        x = x0 + col * (module_w + 28)
        y = 485 + row * 122
        fill = "#E7F4F2" if text not in ("促销中心", "分销设置") else "#F2F4F5"
        centered_box(x, y, module_w, module_h, text, fill=fill)
    draw.text((1240, 716), "促销中心、分销设置当前未启用，后续按计划启用",
              font=small_font, fill="#666666")

    draw.line((900, 765, 900, 815), fill="#087F74", width=5)
    draw.polygon([(890, 805), (900, 820), (910, 805)], fill="#087F74")

    # Internal integrations that support the above functions.
    lane_label("内部系统", 835, 145)
    internal_boxes = ["金蝶商品与库存", "金蝶销售订单", "金蝶销售出库", "微信支付", "权限与审计日志"]
    for i, text in enumerate(internal_boxes):
        centered_box(x0 + i * (box_w + gap), 860, box_w, 94, text,
                     fill="#F2F4F5", font=small_font)

    draw.rounded_rectangle((80, 1010, width - 80, 1080), radius=12, fill="#FFF5D8")
    draw.text((110, 1028),
              "所有调用按授权店铺隔离；商品、订单、售后、物流和账单数据均保留平台业务主键及处理日志。",
              font=small_font, fill="#5E4A00")

    path = ASSET_DIR / filename
    img.save(path, quality=95)
    return path


def build_diagrams():
    return [
        (draw_system_architecture("architecture.png"), "图 4-1 系统对接总体架构"),
        (draw_flow("product.png", "商品档案同步流程", [
            ("金蝶", ["同步商品主数据", "读取库存单位", "读取金蝶编码"]),
            ("抖店", ["查询商品列表", "获取 SPU/SKU", "同步素材与状态"]),
            ("本系统", ["编码匹配", "维护换算比例", "合并完整商品档案", "异常人工校正"]),
        ], "同步顺序固定为：金蝶基础商品在先，抖店 SPU/SKU 信息在后。"), "图 4-2 商品档案与单位映射流程"),
        (draw_flow("order.png", "订单履约与待外采流程", [
            ("抖店", ["接入订单", "查询订单详情"]),
            ("本系统", ["生成平台订单", "按映射比例换算", "校验各子订单可发库存", "充足则自动审核"]),
            ("异常分支", ["不足则暂停下推", "人工生成待外采单", "采购补足库存", "重新校验"]),
            ("金蝶", ["推送销售订单", "返回创建结果"]),
        ], "只有全部子订单可发库存均充足时，平台订单才自动审核并推送金蝶。"), "图 4-3 订单履约与待外采流程"),
        (draw_flow("aftersale.png", "售后、退货与服务订单流程", [
            ("抖店", ["同步售后列表", "获取售后详情", "接收退货物流"]),
            ("本系统", ["匹配原订单", "审核售后", "生成退货单", "质检并入仓"]),
            ("责任判断", ["普通售后归档", "运营责任", "生成服务订单", "运营支付"]),
            ("状态回写", ["回写收货入仓", "同步退款状态", "记录审计日志"]),
        ], "服务订单不改变消费者退货入仓；其客户为承担责任的运营人员。"), "图 4-4 售后、退货与服务订单流程"),
        (draw_flow("delivery.png", "发货与物流回传流程", [
            ("金蝶", ["销售订单创建成功", "生成销售出库单", "写入承运商与单号"]),
            ("本系统", ["逐单生成发货单", "每 6 小时轮询", "校验物流信息", "记录回传结果"]),
            ("抖店", ["申请电子面单", "回传发货", "多包裹处理", "物流纠错"]),
        ], "一个金蝶订单对应一张系统发货单；物流回传失败进入重试和人工待办。"), "图 4-5 发货与物流回传流程"),
    ]


def add_function_chapter(doc) -> None:
    doc.add_page_break()
    add_heading(doc, "3 功能详情", 1)
    add_heading(doc, "3.1 功能一览表", 2)
    rows = [(i + 1, *row) for i, row in enumerate(FUNCTION_ROWS)]
    add_table(doc, ["序号", "功能模块", "功能说明", "关联能力"], rows, widths=[1.0, 3.3, 8.7, 4.0], font_size=7.8)

    add_heading(doc, "3.2 会员管理", 2)
    add_para(doc, "会员管理包含用户注册申请、付款提报列表、微信用户列表、注册用户管理和用户等级。系统按授权店铺管理会员身份及关联信息，汇总会员订单、售后和渠道关系，形成店铺维度的会员档案。")
    add_para(doc, "用户注册申请用于接收和审核注册资料；注册用户管理维护已注册用户的基础信息和状态；微信用户列表记录微信公众号用户的绑定关系；用户等级用于维护分级结果；付款提报列表用于记录会员相关付款提报及处理状态。会员数据仅在对应授权店铺范围内使用。")

    add_heading(doc, "3.3 订单管理", 2)
    add_para(doc, "订单管理包含普通订单列表、待外采单列表、服务订单列表、发货单列表、售后单列表和退货单列表，覆盖抖店订单接入、库存校验、金蝶下单、售后退货、责任结算和物流回传。")
    flow = [
        ("1", "普通订单列表", "接入抖店订单并转换为平台普通订单，保留平台订单号、子订单和授权店铺信息。"),
        ("2", "单位换算与库存校验", "按店内码映射管理中的配比，将平台 SKU 销售数量换算为金蝶库存单位数量，并逐子订单校验可发库存。"),
        ("3", "自动审核与下单", "全部子订单可发库存充足时，订单自动审核通过并推送金蝶创建销售订单。"),
        ("4", "待外采单列表", "存在库存不足子订单时暂不推送金蝶；人工按需生成待外采单，采购补货后重新校验并下推。"),
        ("5", "售后单列表与退货单列表", "同步售后申请及详情，关联原订单；退货场景记录退货物流、仓库质检和入仓结果，并回写平台状态。"),
        ("6", "服务订单列表", "运营责任售后完成正常退货入仓后，以责任运营人员为客户生成服务订单。"),
        ("7", "发货单列表", "金蝶订单创建成功后逐单生成发货单，每 6 小时轮询金蝶销售出库单并将物流信息回传抖店。"),
    ]
    add_table(doc, ["步骤", "功能", "处理规则"], flow, widths=[1.4, 4.0, 11.6], font_size=8.5)
    add_para(doc, "售后处理支持仅退款和退货退款。退货退款由退货单记录物流、签收、质检和入仓结果；实际入仓后回写抖店。物流回传支持单包裹、多包裹和物流信息更正，电子面单申请、更新及取消均与发货单关联。")

    add_heading(doc, "3.4 商品管理", 2)
    add_para(doc, "商品管理包含普通商品列表、金蝶商品分类、商品剂型、店内码映射管理、商品分类、品牌列表和商品评价。商品档案先从金蝶同步商品编码、名称、规格、分类、剂型和库存单位，再通过抖店商品 API 同步对应 SPU、SKU、平台规格、销售状态及素材信息，合并形成完整商品档案。")
    add_para(doc, "店内码映射管理用于维护平台 SKU 销售单位与金蝶商品库存单位之间的编码映射和配比关系。订单数量、可发库存校验以及库存同步均使用同一份生效映射；无法自动匹配的商品由人工维护后方可进入自动履约。")
    add_para(doc, "普通商品列表集中展示内部商品与抖店商品的关联结果；金蝶商品分类、商品剂型、商品分类和品牌列表为商品归类和检索提供基础资料；商品评价用于查看和管理商品评价信息。")

    add_heading(doc, "3.5 财务管理", 2)
    add_para(doc, "财务管理同步抖店店铺动账和结算明细，并按平台订单、售后单、精选联盟推广订单及平台费用进行核对。消费者退款状态以抖店售后退款结果为准，与内部服务订单支付分别管理。")
    add_para(doc, "线下支付仅面向服务订单，用于责任运营人员支付服务订单费用；微信公众号支付完成后回写服务订单支付状态，并与微信支付对账结果核对。")

    add_heading(doc, "3.6 促销中心", 2)
    add_para(doc, "促销中心为系统现有一级功能，当前尚未启用。后续启用时，将在取得相应权限并完成业务配置后，用于承载与店铺商品推广相关的促销管理。")

    add_heading(doc, "3.7 分销设置", 2)
    add_para(doc, "分销设置为系统现有一级功能，当前尚未启用。后续启用时，将结合精选联盟推广商品、联盟订单和佣金数据开展分销业务管理。")

    add_heading(doc, "3.8 运营管理", 2)
    add_para(doc, "运营管理包含公告列表、自定义表单、广告管理、文章管理、库存管理和报表统计。公告列表、广告管理和文章管理用于发布及维护运营内容；自定义表单用于配置业务信息采集表单。")
    add_para(doc, "库存管理展示平台 SKU 库存及同步状态，并结合店内码映射管理完成平台销售单位与金蝶库存单位的换算。报表统计汇总授权店铺的订单、库存、售后、联盟推广和账单数据，供运营人员查阅。")

    add_heading(doc, "3.9 微信公众号", 2)
    add_para(doc, "微信公众号主要面向店铺运营人员，提供订单数据和业务待办查阅，并用于责任运营人员支付服务订单费用。公众号仅展示当前用户有权访问的店铺和服务订单信息。")

    add_heading(doc, "3.10 商城设置", 2)
    add_para(doc, "商城设置包含平台设置、页面设计、签到设置、支付设置、商城服务说明、配送方式、门店列表、物流公司、区域设置、打印机设置和消息配置。")
    add_para(doc, "平台设置维护商城基础参数；支付设置维护服务订单微信支付接口信息；配送方式、物流公司、区域设置和打印机设置服务于发货处理；门店列表维护系统现有店铺资料；消息配置用于系统业务消息参数管理。")

    add_heading(doc, "3.11 后台管理", 2)
    add_para(doc, "后台管理负责系统账号、角色权限、接口任务、基础参数、操作日志和异常记录。后台用户登录时调用商户登录风险 API 完成风险校验；抖店接口调用记录店铺、接口方法、业务主键、返回码和处理时间，敏感凭证不写入业务日志。")


def add_picture_page(doc, path, title, caption, width=16.8):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, FONT_CN, 14, True, INK)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run()
    run.add_picture(str(path), width=Cm(width))
    p3 = doc.add_paragraph(style="Caption")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run(caption)
    set_run_font(r, size=14, color="555555")


def add_architecture_and_screens(doc) -> None:
    doc.add_page_break()
    add_heading(doc, "4 系统架构、功能截图", 1)
    add_heading(doc, "4.1 系统架构、流程图", 2)
    add_para(doc, "以下流程图围绕本次抖店 API 申请范围，说明平台、友联巡店系统、金蝶和微信公众号之间的数据流及关键控制点。")
    for image_path, caption in build_diagrams():
        add_picture_page(doc, image_path, caption, "流程以接口幂等、店铺隔离、状态留痕和异常待办为共同控制要求。", width=16.8)

    doc.add_page_break()
    add_heading(doc, "4.2 系统登录主界面截图", 2)
    for filename, title, caption in SCREENSHOTS[:2]:
        add_picture_page(doc, SOURCE_IMAGES / filename, title, caption)

    doc.add_page_break()
    add_heading(doc, "4.3 系统子功能界面截图", 2)
    add_note(doc, "截图说明", "截图取自现有自研系统，用于证明功能已经具备。申请材料正文统一使用“店铺”措辞；软件名称“友联巡店系统”保持软著登记名称不变。")
    for filename, title, caption in SCREENSHOTS[2:]:
        add_picture_page(doc, SOURCE_IMAGES / filename, title, caption)


CLIENT_CODE = r'''// .NET 9 / ASP.NET 9。密钥和令牌由服务端密钥管理组件提供。
public sealed class DouDianOpenApiClient
{
    private const string BaseUrl = "https://openapi-fxg.jinritemai.com";
    private readonly HttpClient _http;
    private readonly IOfficialDouDianSigner _signer;
    private readonly string _appKey;

    public DouDianOpenApiClient(HttpClient http,
        IOfficialDouDianSigner signer, IConfiguration config)
    {
        _http = http;
        _signer = signer; // 使用官方 SDK/调用指南的 HMAC-SHA256 实现
        _appKey = config["DouDian:AppKey"]!;
    }

    public async Task<JsonDocument> CallAsync(
        string path, string method, object request, ShopToken token,
        CancellationToken ct = default)
    {
        var paramJson = JsonSerializer.Serialize(request);
        var timestamp = DateTimeOffset.Now.ToString("yyyy-MM-dd HH:mm:ss");
        var publicParams = new SortedDictionary<string, string>(StringComparer.Ordinal)
        {
            ["app_key"] = _appKey,
            ["method"] = method,
            ["access_token"] = token.AccessToken,
            ["param_json"] = paramJson,
            ["timestamp"] = timestamp,
            ["v"] = "2",
            ["sign_method"] = "hmac-sha256"
        };
        publicParams["sign"] = _signer.Sign(publicParams);

        var url = QueryHelpers.AddQueryString(BaseUrl + path, publicParams);
        using var response = await _http.GetAsync(url, ct);
        var body = await response.Content.ReadAsStringAsync(ct);
        await ApiAudit.WriteAsync(token.ShopId, method,
            response.StatusCode, body, ct); // 日志脱敏，不落密钥和完整敏感字段
        response.EnsureSuccessStatusCode();
        return JsonDocument.Parse(body);
    }
}
'''


API_CODE = r'''public sealed class DouDianMainApiSamples(DouDianOpenApiClient api)
{
    // 店铺 API：查询店铺经营类目
    public Task<JsonDocument> GetShopCategory(ShopToken t) =>
        api.CallAsync("/shop/getShopCategory", "shop.getShopCategory", new { }, t);

    // 素材中心 API：同步上传商品图片
    public Task<JsonDocument> UploadImage(ShopToken t, string url) =>
        api.CallAsync("/material/uploadImageSync", "material.uploadImageSync", new { image_url = url }, t);

    // 商品 API：分页读取商品
    public Task<JsonDocument> ListProducts(ShopToken t, int page) =>
        api.CallAsync("/product/listV2", "product.listV2", new { page, size = 100 }, t);

    // 商品 API：查询商品详情
    public Task<JsonDocument> GetProduct(ShopToken t, long productId) =>
        api.CallAsync("/product/detail", "product.detail", new { product_id = productId }, t);

    // 商品 API：查询 SKU 详情，合并平台销售单位与规格
    public Task<JsonDocument> GetSku(ShopToken t, long skuId) =>
        api.CallAsync("/sku/detail", "sku.detail", new { sku_id = skuId }, t);

    // 库存 API：读取 SKU 库存
    public Task<JsonDocument> GetStock(ShopToken t, long skuId) =>
        api.CallAsync("/sku/stockNum", "sku.stockNum", new { sku_id = skuId }, t);

    // 库存 API：按换算后的平台销售单位同步库存
    public Task<JsonDocument> SyncStock(ShopToken t, long skuId, int stockNum) =>
        api.CallAsync("/sku/syncStock", "sku.syncStock", new { sku_id = skuId, stock_num = stockNum }, t);

    // 订单 API：增量读取订单
    public Task<JsonDocument> SearchOrders(ShopToken t, long begin, long end) =>
        api.CallAsync("/order/searchList", "order.searchList", new { update_time_start = begin, update_time_end = end }, t);

    // 订单 API：获取订单及子订单详情
    public Task<JsonDocument> GetOrder(ShopToken t, string orderId) =>
        api.CallAsync("/order/orderDetail", "order.orderDetail", new { shop_order_id = orderId }, t);

    // 售后退款 API：增量读取售后单
    public Task<JsonDocument> ListAfterSales(ShopToken t, long begin, long end) =>
        api.CallAsync("/afterSale/List", "afterSale.List", new { start_time = begin, end_time = end }, t);

    // 售后退款 API：读取售后详情并关联原订单
    public Task<JsonDocument> GetAfterSale(ShopToken t, string afterSaleId) =>
        api.CallAsync("/afterSale/Detail", "afterSale.Detail", new { after_sale_id = afterSaleId }, t);

    // 物流发货 API：回传金蝶销售出库取得的物流单号
    public Task<JsonDocument> AddLogistics(ShopToken t, string orderId, string company, string trackingNo) =>
        api.CallAsync("/order/logisticsAdd", "order.logisticsAdd",
            new { order_id = orderId, logistics_code = company, logistics_no = trackingNo }, t);

    // 电子面单 API：创建电子面单订单
    public Task<JsonDocument> CreateWaybill(ShopToken t, object waybillRequest) =>
        api.CallAsync("/logistics/newCreateOrder", "logistics.newCreateOrder", waybillRequest, t);

    // 账单 API：查询店铺动账明细
    public Task<JsonDocument> GetAccountItems(ShopToken t, string start, string end) =>
        api.CallAsync("/order/getShopAccountItem", "order.getShopAccountItem", new { start_time = start, end_time = end }, t);

    // 精选联盟 API：查询店铺联盟推广商品
    public Task<JsonDocument> GetAllianceProducts(ShopToken t, int page) =>
        api.CallAsync("/buyin/queryShopAllianceProducts", "buyin.queryShopAllianceProducts", new { page, page_size = 50 }, t);

    // 保险 API：查询订单保险信息
    public Task<JsonDocument> GetInsurance(ShopToken t, string orderId) =>
        api.CallAsync("/order/insurance", "order.insurance", new { order_id = orderId }, t);

    // 商户登录风险 API：后台用户登录风控检查
    public Task<JsonDocument> CheckLoginRisk(ShopToken t, object riskContext) =>
        api.CallAsync("/antispam/userLogin", "antispam.userLogin", riskContext, t);

    // 飞鸽 API：更新或存储客服工单
    public Task<JsonDocument> NotifyPigeonWorkOrder(ShopToken t, object workOrder) =>
        api.CallAsync("/pigeon/workOrderNotify", "pigeon.workOrderNotify", workOrder, t);

    // 短视频 API：查询已投稿视频列表
    public Task<JsonDocument> ListVideos(ShopToken t, int page) =>
        api.CallAsync("/shopVideo/listVideo", "shopVideo.listVideo", new { page, page_size = 20 }, t);

    // 短视频 API：投稿短视频并关联商品
    public Task<JsonDocument> PublishVideo(ShopToken t, object publishRequest) =>
        api.CallAsync("/shopVideo/publishVideo", "shopVideo.publishVideo", publishRequest, t);
}
'''


FLOW_CODE = r'''public sealed class PlatformOrderFulfillmentService
{
    public async Task HandleAsync(DouDianOrder source, CancellationToken ct)
    {
        // 平台订单号作为幂等键，避免重复生成普通订单。
        var order = await _orders.GetOrCreatePlatformOrderAsync(source, ct);
        foreach (var line in order.Lines)
        {
            var map = await _skuMappings.GetEffectiveAsync(order.ShopId, line.PlatformSkuId, ct);
            line.KingdeeQty = line.PlatformQty * map.KingdeeUnitsPerSaleUnit;
            line.AvailableQty = await _inventory.GetAvailableToShipAsync(map.KingdeeItemCode, ct);
        }

        if (order.Lines.All(x => x.AvailableQty >= x.KingdeeQty))
        {
            await _orders.AutoApproveAsync(order.Id, ct);
            await _kingdee.CreateSalesOrderAsync(order, ct);
            return;
        }

        await _orders.HoldForInventoryAsync(order.Id, ct);
        // 不自动外采；由人工选择不足行后生成待外采单。
        await _tasks.CreateProcurementDecisionTaskAsync(order.Id, ct);
    }
}

public sealed class AfterSaleAndDeliveryJobs
{
    public async Task ReceiveReturnAsync(AfterSale sale, ReturnReceipt receipt, CancellationToken ct)
    {
        await _returns.PutQualifiedGoodsBackToStockAsync(receipt, ct);
        await _douDian.ConfirmReturnWarehouseAsync(sale.ShopId, sale.AfterSaleId, ct);
        if (sale.ResponsibleParty == ResponsibleParty.Operator)
            await _serviceOrders.CreateForOperatorAsync(sale.OperatorId, sale, ct);
    }

    // 由调度器按 0 */6 * * * 触发：每 6 小时一次。
    public async Task PollKingdeeOutboundAsync(CancellationToken ct)
    {
        foreach (var outbound in await _kingdee.GetNewSalesOutboundsAsync(ct))
        {
            var delivery = await _deliveries.GetOrCreateByKingdeeOrderAsync(outbound.OrderNo, ct);
            await _douDian.AddLogisticsAsync(delivery.ShopId, delivery.PlatformOrderId,
                outbound.CarrierCode, outbound.TrackingNo, ct);
            await _deliveries.MarkReportedAsync(delivery.Id, ct);
        }
    }
}
'''


def add_code_block(doc, code, font_size=7.4):
    font_size = max(font_size, 14)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F6F7")
    set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(code.strip().splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line if line else " ")
        set_run_font(r, FONT_CODE, font_size, False, "24313A")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_certificate_and_code(doc) -> None:
    doc.add_page_break()
    add_heading(doc, "4.4 自研发软件著作权", 2)
    add_para(doc, "申请主体已提供友联巡店系统 V1.0 的计算机软件著作权登记证书。证书信息与本申请系统名称、著作权人一致。")
    add_table(doc, ["项目", "证书信息"], [
        ("软件名称", "友联巡店系统 V1.0"),
        ("著作权人", "泉州市友联医疗器械有限公司"),
        ("证书号", "软著登字第12760915号"),
        ("登记号", "2024SR0357042"),
        ("开发完成日期", "2022 年 7 月 8 日"),
        ("登记日期", "2024 年 3 月 6 日"),
    ], widths=[4.2, 12.8], font_size=9)
    add_picture_page(doc, CERT, "友联巡店系统软件著作权登记证书", "图 4-6 计算机软件著作权登记证书（申请主体提供）", width=13.8)

    doc.add_page_break()
    add_heading(doc, "4.5 自研发软件源代码信息", 2)
    add_para(doc, "以下为本系统抖店开放平台对接的核心代码节选。代码采用统一客户端、店铺令牌隔离、官方 HMAC-SHA256 签名组件、脱敏审计和幂等业务服务；实际请求字段以获批权限下的官方接口版本为准。")
    add_note(doc, "安全说明", "示例中的 app_key、app_secret、access_token、店铺编号及业务数据均为配置或对象引用，未展示任何生产凭证。签名由官方 SDK/调用指南实现并经官方调试工具校验。")
    add_heading(doc, "4.5.1 统一开放平台客户端", 3)
    add_code_block(doc, CLIENT_CODE, 7.2)
    add_heading(doc, "4.5.2 主要接口调用示例", 3)
    add_code_block(doc, API_CODE, 6.7)
    add_heading(doc, "4.5.3 订单、售后与发货核心流程节选", 3)
    add_code_block(doc, FLOW_CODE, 6.9)


def normalize_all_runs(doc):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name is None:
                            set_run_font(run)


def validate(doc):
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)
    required = [row[0] for row in API_ROWS]
    missing = [name for name in required if name not in text]
    if missing:
        raise RuntimeError(f"Missing API categories: {missing}")
    for category in required:
        code_label = category.split("（", 1)[0]
        if f"// {code_label}" not in text:
            raise RuntimeError(f"Missing source-code sample for category: {category}")
    if "短信 API 不需要、不申请" not in text:
        raise RuntimeError("SMS exclusion statement missing")
    # “门店列表” is the confirmed current UI function name supplied by the user.
    if "门店" in text.replace("门店列表", ""):
        raise RuntimeError("Unexpected legacy term 门店 remains outside the confirmed UI name 门店列表")


def main():
    if not TEMPLATE.exists() or not SOURCE_DOC.exists() or not CERT.exists():
        raise FileNotFoundError("Reference files are missing")
    for filename, _, _ in SCREENSHOTS:
        if not (SOURCE_IMAGES / filename).exists():
            raise FileNotFoundError(SOURCE_IMAGES / filename)

    # The chapter structure follows the official application template. The source
    # package is used as the style/theme carrier because the supplied blank template
    # has no Word style definitions and renders Chinese as missing glyphs.
    doc = Document(str(SOURCE_DOC))
    clear_document(doc)
    configure_styles(doc)
    configure_section(doc.sections[0], cover=True)
    add_header_footer(doc.sections[0])

    add_cover(doc)
    add_contents(doc)
    add_overview(doc)
    add_api_chapter(doc)
    add_function_chapter(doc)
    add_architecture_and_screens(doc)
    add_certificate_and_code(doc)
    normalize_all_runs(doc)
    validate(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "友联巡店系统 V1.0 抖店商家自研后台系统功能说明书"
    doc.core_properties.subject = "抖店开放平台 API 权限申请"
    doc.core_properties.author = "泉州市友联医疗器械有限公司"
    doc.core_properties.comments = "依据商家自研后台系统功能说明书模板编制"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
